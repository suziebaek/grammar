# doc_formatter.py — 생성 결과를 출판 교재·공인 시험 형식의 Word 문서로 정리
# AI 출력의 서식이 문항마다 달라도(ⓐ/(a)/ㄱ, ___(A)___/______, 깨진 표, 마크다운 등)
# 마지막 출력 단계에서 같은 규칙으로 맞춥니다.
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"
LETTERS = "ⓐⓑⓒⓓⓔⓕⓖⓗ"
JAMO = "ㄱㄴㄷㄹㅁㅂㅅㅇ"
# 받침 있는 번호(일·삼·육·칠·팔·십) — 조사 은/이/을/과
HAS_FINAL = set("①③⑥⑦⑧⑩")
BLANK = "__________"
OPTION_WORD_LIMIT = {"H": 12, "E": 15}

FONT_KO = "맑은 고딕"
FONT_EN = "Arial"


# ───────────────────────── 텍스트 정리 ─────────────────────────
def strip_markdown(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = t.replace("**", "")
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = re.sub(r"[ \t]*[✓✔✗✘]", "", t)
    t = re.sub(r"[🚨🔄✅❌🔺✨💡📊🟦⚠️]", "", t)
    return t


def unify_markers(t):
    t = re.sub(r"\(([a-h])\)", lambda m: LETTERS["abcdefgh".index(m.group(1))], t)
    t = re.sub(r"(?m)^[ \t]*([a-h])[.)][ \t]+", lambda m: LETTERS["abcdefgh".index(m.group(1))] + " ", t)
    t = re.sub(r"(?<![가-힣ㄱ-ㅎ])([ㄱㄴㄷㄹㅁㅂㅅㅇ])(?=[.\s,)은는이가을를과와도의]|$)",
               lambda m: LETTERS[JAMO.index(m.group(1))], t)
    t = re.sub(r"(?m)^([ⓐ-ⓗ])\.[ \t]*", r"\1 ", t)
    t = re.sub(r"(?m)^[ \t]*[•●▪◦][ \t]*", "· ", t)
    t = re.sub(r"(?m)^([ⓐ-ⓗ])[ \t]*(?![은는이가을를과와도의 ,.)·~])", r"\1 ", t)               # 줄 맨 앞 표시: 한 칸 띄움
    t = re.sub(r"(?<=[^\n])([ⓐ-ⓗ])[ \t]+(?=[A-Za-z<\"'])", r"\1", t)  # 문장 속 표시: 붙여 씀
    return t


def unify_blanks(t):
    labeled = {}

    def keep(m):
        key = f"\x00{m.group(1)}\x00"
        labeled[key] = f"_____({m.group(1)})_____"
        return key
    t = re.sub(r"_{2,}[ \t]*\(([A-E])\)[ \t]*_{2,}", keep, t)
    t = re.sub(r"\(([A-E])\)[ \t]*_{3,}", keep, t)
    t = re.sub(r"_{3,}", BLANK, t)
    for k, v in labeled.items():
        t = t.replace(k, v)
    return t


def fix_particles(t):
    def repl(m):
        sym, p = m.group(1), m.group(2)
        if sym in HAS_FINAL:
            table = {"은": "은", "는": "은", "이": "이", "가": "이", "을": "을", "를": "을", "과": "과", "와": "과"}
        else:
            table = {"은": "는", "는": "는", "이": "가", "가": "가", "을": "를", "를": "를", "과": "와", "와": "와"}
        return sym + table[p]
    t = re.sub(r"([①-⑩ⓐ-ⓗ])(은|는|을|를|과|와)(?=[\s,.)]|$)", repl, t)
    t = re.sub(r"([①-⑩ⓐ-ⓗ])(이|가)(?=[\s,]|$)", repl, t)
    return t


def normalize(t):
    t = (t or "").replace("\r", "")
    t = strip_markdown(t)
    t = unify_markers(t)
    t = unify_blanks(t)
    t = fix_particles(t)
    t = re.sub(r"[ \t]+\n", "\n", t)
    return t.strip()


def clean_explanation(t):
    t = normalize(t)
    lines = [l.strip() for l in t.split("\n") if l.strip()]
    return " ".join(lines)


# ───────────────────────── 지문 구조 분석 ─────────────────────────
OPT_RE = re.compile(r"^\s*\|?\s*([①-⑤])\s*\|?\s*(.*)$")
HEADER_RE = re.compile(r"^\s*\|?\s*((?:\([A-E]\)[\s\|/　…]*){2,5})\|?\s*$")
SEP_RE = re.compile(r"\s*(?:\||/|……|…|\t|　+|\s{2,})\s*")


def _split_cols(body):
    return [c for c in SEP_RE.split(body.strip().strip("|")) if c.strip()]


def parse_blocks(passage):
    """지문/보기 텍스트를 블록 목록으로: para, box, table, choice_table, options"""
    lines = [l for l in passage.split("\n")]
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or re.fullmatch(r"\|?[\s\|:\-]*\|?", s) and "|" in s:
            i += 1
            continue
        hm = HEADER_RE.match(s)
        if hm:
            labels = re.findall(r"\([A-E]\)", hm.group(1))
            rows, j = [], i + 1
            while j < len(lines):
                sj = lines[j].strip()
                if not sj or re.fullmatch(r"[\s\|:\-]*", sj):
                    j += 1
                    continue
                om = OPT_RE.match(sj)
                if not om:
                    break
                rows.append((om.group(1), _split_cols(om.group(2))))
                j += 1
            if rows and all(len(c) == len(labels) for _, c in rows):
                blocks.append({"type": "choice_table", "labels": labels, "rows": rows})
            else:
                blocks.append({"type": "para", "text": "  ".join(labels)})
                blocks.append({"type": "options", "items": [(n, " …… ".join(c) if len(c) > 1 else " ".join(c)) for n, c in rows]})
            i = j
            continue
        if s.startswith("|") and not OPT_RE.match(s):
            rows, j = [], i
            while j < len(lines) and lines[j].strip().startswith("|"):
                cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\-\s]*", c) for c in cells):
                    rows.append(cells)
                j += 1
            if rows:
                blocks.append({"type": "table", "rows": rows})
            i = j
            continue
        if s in ("<보기>", "〈보기〉", "[보기]", "<보 기>"):
            items, j = [], i + 1
            while j < len(lines) and not OPT_RE.match(lines[j].strip()):
                if lines[j].strip():
                    items.append(lines[j].strip())
                j += 1
            blocks.append({"type": "box", "items": items})
            i = j
            continue
        if re.match(r"^[ⓐ-ⓗ] ", s):
            items, j = [], i
            while j < len(lines) and re.match(r"^\s*[ⓐ-ⓗ] ", lines[j]):
                items.append(lines[j].strip())
                j += 1
            if len(items) >= 3:
                blocks.append({"type": "box", "items": items})
                i = j
                continue
        om = OPT_RE.match(s)
        if om:
            items, j = [], i
            while j < len(lines):
                sj = lines[j].strip()
                m2 = OPT_RE.match(sj)
                if not m2:
                    if sj:
                        break
                    j += 1
                    continue
                items.append((m2.group(1), m2.group(2).strip().strip("|").strip()))
                j += 1
            blocks.append({"type": "options", "items": items})
            i = j
            continue
        blocks.append({"type": "para", "text": s})
        i += 1
    return blocks


# ───────────────────────── Word 출력 도우미 ─────────────────────────
def _set_fonts(run, size=10, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_KO)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_runs(par, text, size=10, bold=False, color=None):
    for part in re.split(r"(<u>.*?</u>)", text):
        if not part:
            continue
        if part.startswith("<u>") and part.endswith("</u>"):
            r = par.add_run(part[3:-4])
            r.underline = True
        else:
            r = par.add_run(part.replace("<u>", "").replace("</u>", ""))
        _set_fonts(r, size, bold, color)


def _para(container, text="", size=10, bold=False, indent=0.0, first=0.0, before=0, after=2,
          align=None, color=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(first)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.25
    if align:
        p.alignment = align
    if text:
        _add_runs(p, text, size, bold, color)
    return p


def _cell_borders(cell, show=True):
    tcPr = cell._element.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if show else "nil")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tcPr.append(borders)


def _first_par(cell):
    p = cell.paragraphs[0]
    return p


def _write_question(doc, num, stem, passage, level):
    p = _para(doc, "", before=8, after=4)
    _add_runs(p, f"{num}. ", bold=True)
    _add_runs(p, stem)
    for b in parse_blocks(passage):
        t = b["type"]
        if t == "para":
            _para(doc, b["text"], indent=0.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif t == "box":
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            _cell_borders(cell, True)
            p0 = _first_par(cell)
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(p0, "<보기>")
            for it in b["items"]:
                _para(cell, it, indent=0.1)
            _para(doc, "", after=0)
        elif t == "table":
            rows = b["rows"]
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.rows[ri].cells[ci]
                    _cell_borders(cell, True)
                    _add_runs(_first_par(cell), r[ci] if ci < len(r) else "", size=9, bold=(ri == 0))
            _para(doc, "", after=0)
        elif t == "choice_table":
            labels, rows = b["labels"], b["rows"]
            tbl = doc.add_table(rows=len(rows) + 1, cols=len(labels) + 1)
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            tbl.autofit = False
            widths = [Cm(1.0)] + [Cm(3.2)] * len(labels)
            for ci, lab in enumerate([""] + labels):
                cell = tbl.rows[0].cells[ci]
                _cell_borders(cell, False)
                pp = _first_par(cell)
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_runs(pp, lab)
            for ri, (n, cols) in enumerate(rows, start=1):
                for ci, val in enumerate([n] + cols):
                    cell = tbl.rows[ri].cells[ci]
                    _cell_borders(cell, False)
                    pp = _first_par(cell)
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
                    pp.paragraph_format.space_after = Pt(0)
                    pp.paragraph_format.space_before = Pt(0)
                    _add_runs(pp, val)
            tblPr = tbl._tbl.tblPr
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tblPr.append(layout)
            for ci, w in enumerate(widths):
                tbl.columns[ci].width = w
            for row in tbl.rows:
                for ci, cell in enumerate(row.cells):
                    cell.width = widths[ci]
                    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            _para(doc, "", after=0)
        elif t == "options":
            for n, txt in b["items"]:
                _para(doc, f"{n} {txt}", indent=0.6, first=-0.5, after=1)


def _shade(cell, fill="F2F2F2"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _write_answer_box(doc, answer, explanation, diff=""):
    # 🚀 [v6-13] 문항 바로 아래 정답·해설 상자
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _cell_borders(cell, False)
    _shade(cell)
    p0 = _first_par(cell)
    p0.paragraph_format.space_after = Pt(1)
    _add_runs(p0, f"정답 {answer}", size=9, bold=True)
    if diff:
        _add_runs(p0, f"   [{diff}]", size=8, color=(0x59, 0x59, 0x59))
    if explanation:
        p1 = _para(cell, "", after=0)
        _add_runs(p1, "해설  ", size=9, bold=True)
        _add_runs(p1, explanation, size=9)
    _para(doc, "", after=2)


def _option_notes(passage, level):
    notes = []
    limit = OPTION_WORD_LIMIT.get(level, 15)
    for b in parse_blocks(passage):
        if b["type"] == "options":
            if len(b["items"]) != 5:
                notes.append(f"선택지가 {len(b['items'])}개입니다(5개여야 함).")
            for n, txt in b["items"]:
                words = len(re.findall(r"[A-Za-z']+", re.sub(r"<.*?>", "", txt)))
                if words > limit:
                    notes.append(f"선택지 {n}{'이' if n in HAS_FINAL else '가'} {words}단어로 길이 기준({limit}단어)을 넘습니다.")
    return notes


def _collect(entry, level):
    """결과 항목을 (발문, 지문, 정답, 해설, 검토메모)로 정리"""
    out = []
    for r in entry.get("results", []):
        parts = r.get("parsed_parts") or {}
        stem = r.get("clean_question") or parts.get("발문", "")
        passage = r.get("clean_passage") or parts.get("보기/지문", "")
        ans = r.get("clean_answer") or parts.get("정답", "")
        exp = r.get("clean_explanation") or parts.get("해설", "")
        notes = []
        raw_all = "\n".join([str(stem), str(passage)])
        for m in re.findall(r"🚨[^\n]*", raw_all + "\n" + str(r.get("text", ""))):
            notes.append(re.sub(r"[🚨*]", "", m).strip())
        drop = lambda s: "\n".join(l for l in str(s).split("\n")
                                   if not l.strip().startswith("🚨") and not re.match(r"^\s*문제\s*\d+\.?\s*$", l))
        if not str(stem).strip() and not str(passage).strip() and r.get("text"):
            # 🚀 [v6-13] 발문·지문 자동 분리 실패 → 원문(정답 앞부분)을 그대로 사용
            raw = re.split(r"\[\s*정답\s*\]|정답\s*[:：]", str(r["text"]))[0]
            raw = re.sub(r"【문제\s*\d+】|\[\s*(발문|보기\s*/\s*지문|보기|지문|선택지)\s*\]", "", raw).strip()
            _l = raw.split("\n", 1)
            stem, passage = _l[0], (_l[1] if len(_l) > 1 else "")
            notes.append("발문·지문을 자동으로 나누지 못해 원문을 그대로 넣었습니다. 형식을 확인해 주세요.")
        stem, passage = normalize(drop(stem)), normalize(drop(passage))
        stem = re.sub(r"^\s*\d+\.\s*", "", stem).replace("\n", " ")
        ans, exp = clean_explanation(ans), clean_explanation(exp)
        a, c = r.get("assigned_diff"), r.get("raw_diff")
        if a and c and a != c:
            notes.append(f"난이도 차이: 배정 {a} → 실제 {c}")
        if not r.get("is_valid", True):
            notes.append(str(r.get("feedback", "재생성 문항 — 내용 확인 필요")))
        if "밑줄" in stem and "<u>" not in passage:
            notes.append("발문에 '밑줄'이 있으나 지문에 밑줄 표시가 없습니다.")
        notes += _option_notes(passage, level)
        out.append({"stem": stem, "passage": passage, "answer": ans, "explanation": exp,
                    "notes": notes, "diff": c or a or ""})
    return out


def build_exam_docx(history_data, is_multiple=False, level="H", answer_layout="inline"):
    sets = history_data if is_multiple else [history_data]
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    all_items = []
    for si, entry in enumerate(sets, start=1):
        lv = entry.get("level", level)
        items = _collect(entry, lv)
        title = f"{entry.get('major','')} > {entry.get('mid','')}"
        _para(doc, (f"[세트 {si}] " if is_multiple else "") + title, size=12, bold=True, after=2)
        _para(doc, f"출제 개념: {entry.get('minor','')}  |  {entry.get('difficulty','')}", size=9,
              color=(0x59, 0x59, 0x59), after=6)
        for qi, it in enumerate(items, start=1):
            _write_question(doc, qi, it["stem"], it["passage"], lv)
            if answer_layout == "inline":
                _write_answer_box(doc, it["answer"], it["explanation"], it["diff"])
        all_items.append((si, title, items))

    # 정답 및 해설 (뒤쪽에 모으는 방식: answer_layout="end"일 때만)
    if answer_layout == "end":
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _para(doc, "정답 및 해설", size=12, bold=True, after=6)
        for si, title, items in all_items:
            if is_multiple:
                _para(doc, f"[세트 {si}] {title}", size=10, bold=True, before=6)
            for qi, it in enumerate(items, start=1):
                p = _para(doc, "", before=4, after=1)
                _add_runs(p, f"{qi}. 정답 {it['answer']}", bold=True)
                if it["diff"]:
                    _add_runs(p, f"  [{it['diff']}]", size=9, color=(0x59, 0x59, 0x59))
                if it["explanation"]:
                    _para(doc, it["explanation"], indent=0.4, after=2)

    # 검토 메모(출판 전 확인용)
    noted = [(si, qi, it["notes"]) for si, _, items in all_items for qi, it in enumerate(items, start=1) if it["notes"]]
    if noted:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _para(doc, "검토 메모 (출판 전 확인 후 삭제)", size=11, bold=True, after=6, color=(0xC0, 0x00, 0x00))
        for si, qi, notes in noted:
            head = f"[세트 {si}] {qi}번" if is_multiple else f"{qi}번"
            for n in notes:
                _para(doc, f"· {head}: {n}", size=9, indent=0.2, after=1)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
